import json,re,sys
sys.path.insert(0,'.')
from docx import Document
from docx.shared import Pt,RGBColor,Cm
from bib_over import ALIAS
d=json.load(open("refs_out.json")); REF=d["R"]; INC=d["I"]
s2=json.load(open("sinref2.json")); res=s2["res"]; cat=s2["cat"]
doc=Document()
st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10.5)
for sec in doc.sections: sec.left_margin=sec.right_margin=Cm(2.2)
doc.add_heading("Referencias pendientes — Psiconautas",1)
doc.add_paragraph("Generado el 2026-10-15 a partir de src/datos/bibliografia.js. Parte 1: citas de la app que no se pudieron enlazar con ninguna referencia. Parte 2: referencias que existen pero les falta algún dato.")
doc.add_heading(f"Parte 1. Citas sin referencia resuelta ({sum(len(v) for v in cat.values())})",2)
etq={"sin_autor_año":"A. Sin autor o sin año identificable (documentos de cátedra, obras completas, «s. f.») — {n}","con_autor_año":"B. Con autor y año, pero sin entrada en el registro de referencias — {n}"}
for k in ("con_autor_año","sin_autor_año"):
    doc.add_heading(etq[k].format(n=len(cat[k])),3)
    for b in cat[k]:
        f,t=res.get(b,("?",""))
        p=doc.add_paragraph(style="List Bullet"); p.add_run(b)
        r=p.add_run(f"\n   Se usa en: {f}"+(f" → «{t}»" if t else "")); r.italic=True; r.font.size=Pt(9)
doc.add_heading("Parte 2. Referencias incompletas",2)
filas=[]
for k in INC:
    t=REF.get(k)
    al=None
    if t is None:
        al=ALIAS.get(k); t=REF.get(al,"")
    falta=[]
    if "(s. f.)" in t: falta.append("año")
    if not re.search(r"\.\s+(?:\*[^*]+\*|[A-ZÁÉÍÓÚ][^.]*)\.\s+[A-ZÁÉÍÓÚ]",t) and not re.search(r"https://|doi",t): falta.append("editorial o fuente")
    if re.search(r"\*[^*]*,\s*en\s",t): falta.append("el título está mezclado con «en …»")
    if "En " not in t and not re.search(r"\d+\s*\(?\d*\)?,\s*[\d–]+",t) and re.search(r"Journal|Review|Bulletin|Psychologist|Revista",t): falta.append("volumen y páginas")
    filas.append((k,t,falta or ["editorial, páginas o datos de edición"],al))
for k,t,falta,al in filas:
    p=doc.add_paragraph(style="List Number"); p.add_run(t or k)
    r=p.add_run("\n   Falta: "+"; ".join(falta)+(f" (la clave {k} se resuelve a otra entrada ya registrada)" if al else "")); r.italic=True; r.font.size=Pt(9)
doc.save("/Users/nicolasayasandoval/Downloads/psicoterapias-app/auditoria_notebooklm/docx/Referencias_pendientes.docx")
print(len(filas))
for k,t,f,a in filas[:8]: print(k,f)
